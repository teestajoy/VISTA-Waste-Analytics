"""JSON / CSV / Excel exports generated from the stored report (single source of truth)."""
import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def flat_row(r: dict) -> dict:
    c, d, cont, cond, imp, ex = r["classification"], r["detection"], r["contamination"], r["condition"], r["carbon_economic"], r["explainability"]
    row = {
        "report_id": r["report_id"], "timestamp_utc": r["timestamp_utc"], "schema_version": r["report_schema_version"],
        "filename": r["input"]["filename"], "image_width": r["input"]["width"], "image_height": r["input"]["height"],
        "detection_id": d["detection_id"], "detector_label": d["label"], "detector_confidence": d["confidence"],
        "box_x0": d["box_xyxy"][0], "box_y0": d["box_xyxy"][1], "box_x1": d["box_xyxy"][2], "box_y1": d["box_xyxy"][3],
        "stream": c["stream"]["prediction"], "stream_confidence": c["stream"]["confidence"], "stream_source": c["stream"]["source"],
        "category": c["fine"]["prediction"], "category_confidence": c["fine"]["confidence"], "uncertain": c["uncertain"],
        "contamination_proxy": cont.get("score"), "saturation_variance": cont.get("saturation_variance"),
        "condition_score": cond.get("score"), "solidity": cond.get("solidity"), "contour_area": cond.get("contour_area"),
        "convex_hull_area": cond.get("convex_hull_area"),
        "mass_kg": imp.get("mass_kg"), "mass_method": imp.get("mass_method"), "co2e_avoided_kg": imp.get("co2e_avoided_kg"),
        "recovery_value_usd": imp.get("recovery_value_usd"), "factors_are_placeholders": imp.get("factors_are_placeholders"),
        "gradcam_status": ex["status"], "gradcam_target": ex["target_class"], "gradcam_layer": ex["target_layer"],
        "total_ms": r["timings"]["total_ms"], "warnings": " | ".join(r["warnings"]),
    }
    for k, v in c["fine"]["probabilities"].items():
        row[f"prob_{k}"] = v
    return row


def to_json(r: dict) -> bytes:
    return json.dumps(r, indent=2, ensure_ascii=False).encode()


def to_csv(rows: list[dict]) -> bytes:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in cols})
    return buf.getvalue().encode("utf-8-sig")


def _sheet(wb: Workbook, name: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
    for r in rows:
        ws.append(["N/A" if v is None else v for v in r])
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"


def to_xlsx(r: dict, evaluation: dict | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    f = flat_row(r)
    _sheet(wb, "Analysis Summary", ["Field", "Value"], [[k, v] for k, v in f.items() if not k.startswith("prob_")])
    d = r["detection"]
    _sheet(wb, "Detection Results", ["Detection ID", "Label", "Confidence", "x0", "y0", "x1", "y1", "Crop xyxy"],
           [[d["detection_id"], d["label"], d["confidence"], *d["box_xyxy"], str(d["crop_xyxy"])]])
    c = r["classification"]
    _sheet(wb, "Classification Probabilities", ["Level", "Class", "Probability"],
           [["stream", k, v] for k, v in c["stream"]["probabilities"].items()] + [["fine", k, v] for k, v in c["fine"]["probabilities"].items()])
    cont, cond = r["contamination"], r["condition"]
    _sheet(wb, "Contamination and Condition", ["Measure", "Value", "Note"], [
        ["Contamination proxy score", cont.get("score"), cont["label"]], ["Saturation variance", cont.get("saturation_variance"), ""],
        ["Laplacian energy", cont.get("laplacian_energy"), ""], ["Contamination formula", cont["formula"], ""],
        ["Condition score", cond.get("score"), cond["label"]], ["Solidity", cond.get("solidity"), ""],
        ["Contour area", cond.get("contour_area"), ""], ["Convex hull area", cond.get("convex_hull_area"), ""], ["Condition formula", cond["formula"], ""]])
    imp = r["carbon_economic"]
    _sheet(wb, "Carbon and Recovery Value", ["Item", "Value"], [[k, v] for k, v in imp.items()])
    ex = r["explainability"]
    _sheet(wb, "Explainability Metadata", ["Item", "Value"], [[k, json.dumps(v) if isinstance(v, (dict, list)) else v] for k, v in ex.items()])
    if evaluation:
        _sheet(wb, "Model Evaluation", ["Model", "Metric", "Value"],
               [[m, k, v] for m, e in evaluation.items() for k, v in (e.get("metrics") or {}).items() if isinstance(v, (int, float))])
    _sheet(wb, "Reference Factors and Sources", ["Item", "Value"], [["source", imp.get("source")], ["geographic_scope", imp.get("geographic_scope")],
                                                                     ["reference_year", imp.get("reference_year")], ["notice", imp.get("notice")]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
