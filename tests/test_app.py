"""Tests for the web application back end. A tiny random-weight ResNet18 is used ONLY to test plumbing (shapes, routing, exports)."""
import io, json
import numpy as np, pytest, torch
from PIL import Image
from torchvision import models

from vista.app import exports, explain, hier, measure
from vista.app.models import PREP, Loaded, ModelManager

CLASSES = ["battery", "cardboard", "glass", "metal", "organic", "paper", "plastic", "textile", "trash"]


def _mgr(stream=False):
    torch.manual_seed(0)
    m = ModelManager()
    m.fine = Loaded("fine", "x", models.resnet18(weights=None, num_classes=9).eval(), CLASSES)
    if stream:
        m.stream = Loaded("stream", "x", models.resnet18(weights=None, num_classes=2).eval(), ["MSW", "E_WASTE"])
    return m


def _img(seed=0, size=96):
    return Image.fromarray(np.random.default_rng(seed).integers(0, 255, (size, size, 3), dtype=np.uint8))


def test_routing_restricts_labels_to_stream():
    for seed in range(6):
        r = hier.classify(_img(seed), _mgr(), 0.7, 0.6); r.pop("_ctx")
        fam = r["stream"]["prediction"]
        assert hier.family_of(r["fine"]["raw_label"]) == fam
        assert abs(sum(r["fine"]["probabilities"].values()) - 1) < 1e-3
        assert abs(sum(r["stream"]["probabilities"].values()) - 1) < 1e-3


def test_rule_gate_uses_detector_label():
    r = hier.classify(_img(), _mgr(), 0.7, 0.6, det_label="cell phone"); r.pop("_ctx")
    assert r["stream"]["prediction"] == "E_WASTE" and r["stream"]["source"] == "rule_based_detector_label"
    assert hier.family_of(r["fine"]["raw_label"]) == "E_WASTE"


def test_dedicated_stream_model_is_used():
    r = hier.classify(_img(), _mgr(stream=True), 0.7, 0.6); r.pop("_ctx")
    assert r["stream"]["source"] == "dedicated_stream_model"


def test_uncertain_when_below_threshold():
    r = hier.classify(_img(), _mgr(), 0.999, 0.999); r.pop("_ctx")
    assert r["uncertain"] and r["uncertainty_reasons"]


def test_no_classifier_raises():
    with pytest.raises(hier.ClassificationUnavailable):
        hier.classify(_img(), ModelManager(), 0.7, 0.6)


def test_gradcam_shape_range_target_and_hooks():
    m = models.resnet18(weights=None, num_classes=9).eval(); x = PREP(_img()).unsqueeze(0)
    a = explain.gradcam(m, x, 0); b = explain.gradcam(m, x, 3)
    assert a["cam"].shape == (224, 224) and np.isfinite(a["cam"]).all() and 0 <= a["cam"].min() and a["cam"].max() <= 1
    assert not np.allclose(a["cam"], b["cam"])                       # map changes with the target class
    assert len(m.layer4._forward_hooks) == 0                         # hooks removed
    assert all(p.grad is None for p in m.parameters())               # no gradient accumulation


def test_gradcam_degenerate_raises():
    m = models.resnet18(weights=None, num_classes=9).eval()
    torch.nn.init.zeros_(m.fc.weight)
    with pytest.raises(explain.ExplainError):
        explain.gradcam(m, PREP(_img()).unsqueeze(0), 0)


def test_measure_edge_cases():
    flat = Image.new("RGB", (64, 64), (120, 120, 120))
    c = measure.contamination(flat); assert 0 <= c["score"] <= 1 and c["saturation_variance"] == 0
    cond, _ = measure.condition(flat); assert cond["score"] is None or 0 <= cond["score"] <= 1
    tiny, _ = measure.condition(Image.new("RGB", (4, 4))); assert tiny["score"] is None
    im = Image.new("RGB", (100, 100), "white")
    from PIL import ImageDraw; ImageDraw.Draw(im).ellipse([20, 20, 80, 80], fill="black")
    cond, ov = measure.condition(im); assert cond["solidity"] > 0.9 and ov is not None


def test_impact_formula_and_unavailable():
    r = measure.impact("plastic", 2.0); assert r["co2e_avoided_kg"] == round(r["co2e_factor_kg_per_kg"] * 2.0, 5) and r["factors_are_placeholders"]
    assert measure.impact("keyboard", None)["status"] == "Not available"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    import vista.app.config as cfg
    monkeypatch.setattr(cfg, "DATA_DIR", tmp_path); monkeypatch.setattr(cfg, "UPLOAD_DIR", tmp_path / "up"); monkeypatch.setattr(cfg, "RUN_DIR", tmp_path / "runs")
    monkeypatch.setattr(cfg, "DB_PATH", tmp_path / "db.sqlite"); monkeypatch.setattr(cfg, "EVAL_DIR", tmp_path / "ev")
    from fastapi.testclient import TestClient
    from vista.app import api
    from vista.app.models import manager
    class FakeYolo:
        def predict(self, img, **kw):
            import types
            box = types.SimpleNamespace(xyxy=[torch.tensor([10., 10., 80., 80.])], cls=torch.tensor(0), conf=torch.tensor(0.9))
            return [types.SimpleNamespace(boxes=[box], names={0: "bottle"})]
    f = _mgr(); manager.fine, manager.stream, manager.yolo = f.fine, None, FakeYolo()
    monkeypatch.setattr(api.manager, "load_all", lambda: None)
    with TestClient(api.app) as c:
        yield c


def _png():
    b = io.BytesIO(); _img(1, 128).save(b, "PNG"); return b.getvalue()


def test_full_flow_and_exports(client):
    d = client.post("/api/detect", files={"file": ("a.png", _png(), "image/png")}).json()
    assert d["detections"][0]["label"] == "bottle"
    rep = client.post("/api/analyze", json={"image_id": d["image_id"], "detection_id": 0, "mass_kg": 0.5}).json()
    aid = rep["report_id"]; ce = rep["carbon_economic"]; assert ce["status"] == "Not available" or ce["mass_method"] == "user-entered mass"
    j = json.loads(client.get(f"/api/reports/{aid}/json").content); assert j["report_id"] == aid
    csv_rows = client.get(f"/api/reports/{aid}/csv").content.decode("utf-8-sig").splitlines(); assert len(csv_rows) == 2 and aid in csv_rows[1]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(f"/api/reports/{aid}/excel").content)); assert "Analysis Summary" in wb.sheetnames and "Reference Factors and Sources" in wb.sheetnames
    for k in ("annotated-image", "gradcam", "heatmap", "crop"): assert client.get(f"/api/reports/{aid}/{k}").status_code == 200
    e = client.post("/api/explain/gradcam", json={"analysis_id": aid, "target": "predicted"}); assert e.status_code == 200
    assert client.post("/api/explain/gradcam", json={"analysis_id": aid, "target": "stream"}).status_code == 422   # no dedicated stream model
    ov = client.get("/api/analytics/overview").json(); assert ov["total_analyses"] == 1
    assert client.get("/api/analytics/model-performance").json()["empty"] is True
    assert client.get("/api/history", params={"stream": "NOPE"}).json() == []


def test_validation_errors(client):
    assert client.post("/api/detect", files={"file": ("a.gif", b"x", "image/gif")}).status_code == 422
    assert client.post("/api/detect", files={"file": ("a.png", b"notanimage", "image/png")}).status_code == 422
    assert client.post("/api/analyze", json={"image_id": "nope"}).status_code == 422
    assert client.get("/api/history/zzz").status_code == 404


def test_empty_analytics(client):
    assert client.get("/api/analytics/overview").json()["empty"] is True
